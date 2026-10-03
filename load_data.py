from pathlib import Path
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

DATA_DIR = Path("data")


def read_word(path):
    """Read a Word file and keep paragraphs and tables in their original order."""
    doc = Document(path)
    lines = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = Paragraph(child, doc).text.strip()
            if text:
                lines.append(text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                cells = [c.text.strip() for c in row.cells]
                lines.append(" | ".join(cells))
    return lines


def read_excel(path):
    """Read every sheet of an Excel file, one text line per non-empty row."""
    wb = load_workbook(path, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                cells = ["" if v is None else str(v) for v in row]
                lines.append(f"[{ws.title}] " + " | ".join(cells))
    return lines


if __name__ == "__main__":
    word_lines = read_word(DATA_DIR / "apple_valuation_report.docx")
    excel_lines = read_excel(DATA_DIR / "apple_financial_statements.xlsx")

    print("Word lines:", len(word_lines))
    print("Excel lines:", len(excel_lines))
    print("--- First 5 Word lines ---")
    for line in word_lines[:5]:
        print(line)
    print("--- First 5 Excel lines ---")
    for line in excel_lines[:5]:
        print(line)